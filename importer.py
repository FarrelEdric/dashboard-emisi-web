"""Isi database: data awal dari data.json, impor dari Excel (.xlsx) atau CSV.

Lewat terminal:  python importer.py data.xlsx --year 2026 --quarter 2
Lewat website:   halaman Admin, bagian "Unggah data".
"""
import csv
import io
import json
import re
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import and_, or_

from models import Airport, Route, RouteFlight, db

ALIAS = {"BLI": "DPS"}  # di Excel Bali ditulis BLI
PATTERN = re.compile(r"\b([A-Z]{3})\s*-\s*([A-Z]{3})\b")
TEMPLATE_CSV = ("rute,hemat_nm,hemat_menit,penerbangan_per_hari\n"
                "CGK-DPS,12.1,1.7,50\n"
                "CGK-SUB,14.3,2.0,44\n")

def seed_if_empty():
    """Jika database kosong, isi dengan data contoh dari data.json (tahun 2025 dan 2026)."""
    if Airport.query.first():
        return
    try:
        d = json.loads((Path(__file__).parent / "data.json").read_text(encoding="utf-8"))
        for code, (name, lat, lon, dx, dy, anchor) in d["airports"].items():
            db.session.add(Airport(code=code, name=name, lat=lat, lon=lon, dx=dx, dy=dy, anchor=anchor))
        db.session.flush()
        for a, b, nm, mn, fl in d["routes"]:
            r = Route(origin=a, dest=b, saving_nm=nm, saving_min=mn)
            db.session.add(r)
            db.session.flush()
            for year in (2025, 2026):
                for q in range(1, 5):
                    db.session.add(RouteFlight(route_id=r.id, year=year, quarter=q, flights_per_day=fl[q - 1]))
        db.session.commit()
    except Exception:
        # Diisi lebih dulu oleh proses lain yang jalan bersamaan (umum di serverless). Aman diabaikan.
        db.session.rollback()


def _save_row(a, b, nm, minutes, flights, year, quarter):
    """Simpan satu rute + penerbangan per hari (tambah baru, atau perbarui jika sudah ada)."""
    a, b = ALIAS.get(a, a), ALIAS.get(b, b)
    for code in (a, b):
        if not db.session.get(Airport, code):
            raise ValueError(f"Bandara {code} belum ada. Tambahkan dulu di bagian Bandara.")
    route = Route.query.filter(or_(and_(Route.origin == a, Route.dest == b),
                                   and_(Route.origin == b, Route.dest == a))).first()
    if route:
        route.saving_nm, route.saving_min = nm, minutes
    else:
        route = Route(origin=a, dest=b, saving_nm=nm, saving_min=minutes)
        db.session.add(route)
    db.session.flush()
    rf = db.session.get(RouteFlight, (route.id, year, quarter))
    if rf:
        rf.flights_per_day = flights
    else:
        db.session.add(RouteFlight(route_id=route.id, year=year, quarter=quarter, flights_per_day=flights))


def import_excel(source, year, quarter):
    """Excel RNAV: kolom A rute (mis. CGK - SUB), C hemat Nm, E hemat menit, F penerbangan/hari.
    `source` boleh nama file atau objek file (BytesIO)."""
    ws = load_workbook(source, data_only=True).active
    count = 0
    for row in ws.iter_rows(values_only=True):
        m = PATTERN.search(str(row[0] or "").upper())
        if not m:
            continue
        try:
            nm, minutes, flights = float(row[2]), float(row[4]), float(row[5])
        except (TypeError, ValueError, IndexError):
            continue
        _save_row(m.group(1), m.group(2), nm, minutes, flights, year, quarter)
        count += 1
    db.session.commit()
    return count


def _num(value, label, line):
    try:
        return float(str(value).strip().replace(",", "."))
    except ValueError:
        raise ValueError(f"Baris {line}: kolom {label} harus berupa angka.")


def import_csv(text, year, quarter):
    """CSV dengan header: rute,hemat_nm,hemat_menit,penerbangan_per_hari. Pemisah koma atau titik koma."""
    first = text.splitlines()[0] if text.strip() else ""
    delimiter = ";" if first.count(";") > first.count(",") else ","
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    reader.fieldnames = [(h or "").strip().lower() for h in (reader.fieldnames or [])]
    need = {"rute", "hemat_nm", "hemat_menit", "penerbangan_per_hari"}
    if not need.issubset(reader.fieldnames):
        raise ValueError("Header CSV harus: rute, hemat_nm, hemat_menit, penerbangan_per_hari.")
    count = 0
    for line, row in enumerate(reader, start=2):
        if not any((v or "").strip() for v in row.values() if isinstance(v, str)):
            continue
        m = PATTERN.search(str(row["rute"] or "").upper())
        if not m:
            raise ValueError(f"Baris {line}: rute harus berbentuk CGK-DPS.")
        _save_row(m.group(1), m.group(2), _num(row["hemat_nm"], "hemat_nm", line),
                  _num(row["hemat_menit"], "hemat_menit", line),
                  _num(row["penerbangan_per_hari"], "penerbangan_per_hari", line), year, quarter)
        count += 1
    db.session.commit()
    return count


if __name__ == "__main__":
    import argparse
    from app import app

    p = argparse.ArgumentParser()
    p.add_argument("file")
    p.add_argument("--year", type=int, required=True)
    p.add_argument("--quarter", type=int, required=True, choices=[1, 2, 3, 4])
    a = p.parse_args()
    with app.app_context():
        try:
            if a.file.lower().endswith(".csv"):
                n = import_csv(Path(a.file).read_text(encoding="utf-8-sig"), a.year, a.quarter)
            else:
                n = import_excel(a.file, a.year, a.quarter)
            print(n, "rute berhasil diimpor")
        except ValueError as e:
            db.session.rollback()
            print("Gagal:", e)
